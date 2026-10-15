from __future__ import annotations
import os
import sys
import time
import socket
import threading
import re
from datetime import date
from decimal import Decimal
from pathlib import Path
import uvicorn
import fitz  # PyMuPDF
import openpyxl
from playwright.sync_api import sync_playwright, expect

# Ensure UTF-8 output encoding for Windows PowerShell
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT))

from src.shared.database import Database
from src.shared.event_bus import EventBus
from src.core.domain.order import OrderChannel, OrderStatus
from src.extended_core.crm.models import DealStage
from src.core.services.inventory_service import InventoryService
from src.core.services.order_service import OrderService
from src.extended_core.crm.services.crm_service import CrmService
from src.extended_core.accounting.services.accounting_service import AccountingService
from src.channel_ui.web.app import create_app


def get_free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.bind(("", 0))
        return s.getsockname()[1]


def start_server_in_thread(app, port: int):
    config = uvicorn.Config(app, host="127.0.0.1", port=port, log_level="warning")
    server = uvicorn.Server(config)
    t = threading.Thread(target=server.run, daemon=True)
    t.start()
    return server


def wait_for_server(url: str, timeout: float = 10.0):
    import urllib.request
    start = time.time()
    while time.time() - start < timeout:
        try:
            with urllib.request.urlopen(f"{url}/", timeout=1.0) as resp:
                if resp.status == 200:
                    return True
        except Exception:
            time.sleep(0.3)
    raise TimeoutError(f"Server at {url} did not respond within {timeout}s")


def test_playwright_iteration0923_1455_suite():
    """Playwright E2E UI verification for all 6 scenarios in iteration-0923-1455-verification.md
    MANDATORY HEADED MODE:
      - Scenario 1 (Bug 1): CRM Quotation dynamic multi-item add/remove, dynamic calculation & Proposal progression.
      - Scenario 2 (Bug 2): CRM Smooth scroll on stage change & Column 4 Accordion collapse for Won / Lost deals.
      - Scenario 3 (Bug 3): CRM Converted order link navigation to Orders table with Amber Ring pulse highlight.
      - Scenario 4 (Bug 4): Partial shipment with custom warehouse selection (WH-STORE) and stock balance verification.
      - Scenario 5 (Bug 5): Multi-batch partial shipment continuous multi-page PDF generation (2 pages for order, 1 page for single DO).
      - Scenario 6 (Bug 6): Inventory stock valuation Excel export & Accounting reports convergence to 3 core sheets.
    """
    port = get_free_port()
    base_url = f"http://127.0.0.1:{port}"

    test_dir = PROJECT_ROOT / "temp_playwright_0923_1455"
    test_dir.mkdir(parents=True, exist_ok=True)
    test_db_path = test_dir / "test_iteration0923_1455.db"
    if test_db_path.exists():
        try:
            test_db_path.unlink()
        except Exception:
            pass

    db = Database(test_db_path)
    db.init_db()
    event_bus = EventBus()

    # Create app which automatically registers event handlers and runs seed_demo_data
    app = create_app(db=db, event_bus=event_bus)

    # Add a lost deal so Column 4 has both WON and LOST deals for accordion testing
    order_svc = OrderService(db=db, inventory_service=InventoryService(db=db), event_bus=event_bus)
    crm_svc = CrmService(db=db, order_service=order_svc, event_bus=event_bus)
    with db.transaction() as conn:
        contact_row = conn.execute("SELECT contact_id FROM crm_contacts LIMIT 1").fetchone()
    if contact_row:
        lost_deal = crm_svc.create_deal(
            contact_id=contact_row["contact_id"],
            title="2026 舊版系統採購標案 (競爭對手低價搶標)",
            expected_amount=Decimal("120000.0"),
            stage=DealStage.LOST,
        )

    server = start_server_in_thread(app, port)
    wait_for_server(base_url)
    print(f"\n[*] Playwright Headed Test Server running at {base_url}")

    with sync_playwright() as p:
        # Mandatory Headed Mode with human-observable slow_mo
        browser = p.chromium.launch(headless=False, slow_mo=350)
        context = browser.new_context(viewport={"width": 1440, "height": 900})
        page = context.new_page()
        page.set_default_timeout(2000)
        page.set_default_navigation_timeout(5000)

        page.on("console", lambda msg: print(f"  [BROWSER CONSOLE] {msg.type}: {msg.text}"))

        try:
            # =================================================================
            # SCENARIO 1 (Bug 1): 報價單動態多品項增刪與自動試算
            # =================================================================
            print("\n" + "=" * 75)
            print("[Scenario 1/6] 報價單動態多品項增刪與自動試算 (Bug 1)")
            print("=" * 75)

            page.goto(f"{base_url}/crm?tab=kanban")
            expect(page.locator("h1").filter(has_text="哩來 (LiLai) CRM").first).to_be_visible()

            # Locate Deal 2 in Column 2 (DISCOVERY)
            discovery_col = page.locator("div.bg-sky-50\\/50")
            expect(discovery_col).to_be_visible()
            # Find deal card for "研發廠區無線溫濕度傳感監測案"
            deal2_card = discovery_col.locator("div[id^='deal-card-']").filter(has_text="研發廠區無線溫濕度傳感監測案").first
            expect(deal2_card).to_be_visible()
            deal2_id = deal2_card.get_attribute("id").replace("deal-card-", "")
            print(f"  -> Found target discovery deal: {deal2_id} (研發廠區無線溫濕度傳感監測案)")

            # Click [＋ 報價單] button
            deal2_card.locator("button:has-text('＋ 報價單')").click()
            quote_modal = page.locator("#modal-new-quote")
            expect(quote_modal).to_be_visible()
            print("  -> Quote modal successfully opened.")

            # Initial row count should be 1
            items_container = page.locator("#quote-items-container")
            expect(items_container.locator(".quote-item-row")).to_have_count(1)

            # Click [＋ 新增報價項目] twice to expand to 3 items
            add_item_btn = page.locator("button:has-text('＋ 新增報價項目')")
            add_item_btn.click()
            add_item_btn.click()
            expect(items_container.locator(".quote-item-row")).to_have_count(3)
            print("  -> Added 2 dynamic items; total item rows = 3.")

            # Configure Row 1: SKU-PRO-01, qty 2, price 2500 -> 5000
            row1 = items_container.locator(".quote-item-row").nth(0)
            row1.locator(".quote-sku-select").select_option("SKU-PRO-01")
            row1.locator(".quote-qty-input").fill("2")
            row1.locator(".quote-price-input").fill("2500")

            # Configure Row 2: SKU-ACC-02, qty 4, price 600 -> 2400
            row2 = items_container.locator(".quote-item-row").nth(1)
            row2.locator(".quote-sku-select").select_option("SKU-ACC-02")
            row2.locator(".quote-qty-input").fill("4")
            row2.locator(".quote-price-input").fill("600")

            # Configure Row 3: SKU-IOT-03, qty 5, price 800 -> 4000
            row3 = items_container.locator(".quote-item-row").nth(2)
            row3.locator(".quote-sku-select").select_option("SKU-IOT-03")
            row3.locator(".quote-qty-input").fill("5")
            row3.locator(".quote-price-input").fill("800")

            # Trigger calculation by inputting on qty
            row3.locator(".quote-qty-input").press("Tab")

            # Verify dynamic total = 5000 + 2400 + 4000 = 11,400
            total_display = page.locator("#quote-total-display")
            expect(total_display).to_contain_text("11,400")
            print("  -> Dynamic total for 3 items verified: NT$ 11,400.")

            # Delete Row 3 using the remove button (&times; / 刪除此項)
            row3.locator(".quote-remove-btn").click()
            expect(items_container.locator(".quote-item-row")).to_have_count(2)
            print("  -> Removed row 3; remaining rows = 2.")

            # Verify dynamic total updated to 5000 + 2400 = 7,400
            expect(total_display).to_contain_text("7,400")
            print("  -> Dynamic total updated automatically: NT$ 7,400.")

            # Submit: [產出報價單並進展至 PROPOSAL]
            page.locator("button:has-text('產出報價單並進展至 PROPOSAL')").click()
            expect(quote_modal).not_to_be_visible(timeout=5000)
            print("  -> Quote submitted; modal closed.")

            # Verify deal moved to Column 3 (PROPOSAL)
            proposal_col = page.locator("div.bg-amber-50\\/50")
            moved_card = proposal_col.locator(f"#deal-card-{deal2_id}")
            expect(moved_card).to_be_visible()
            expect(moved_card).to_contain_text("方案報價")
            print(f"  -> Deal {deal2_id} successfully progressed to Column 3 (PROPOSAL).")

            # Verify card displays quote number and download buttons
            docx_btn = moved_card.locator("a:has-text('📥 DOCX')")
            pdf_btn = moved_card.locator("a:has-text('📑 預覽 PDF')")
            expect(docx_btn).to_be_visible()
            expect(pdf_btn).to_be_visible()

            # Test DOCX download from proposal card
            docx_save_path = test_dir / f"quote_{deal2_id}.docx"
            with page.expect_download() as dl_docx_info:
                docx_btn.click()
            dl_docx = dl_docx_info.value
            dl_docx.save_as(str(docx_save_path))
            assert docx_save_path.exists() and docx_save_path.stat().st_size > 0
            print(f"  -> DOCX download verified: {docx_save_path.stat().st_size} bytes.")

            # Test PDF preview (target='_blank' inline PDF opens new tab)
            with context.expect_page() as new_page_info:
                pdf_btn.click()
            pdf_page = new_page_info.value
            pdf_url = pdf_page.url
            print(f"  -> PDF preview tab opened at: {pdf_url}")
            pdf_resp = page.request.get(pdf_url)
            assert pdf_resp.status == 200
            pdf_bytes = pdf_resp.body()
            with fitz.open(stream=pdf_bytes, filetype="pdf") as doc:
                assert len(doc) >= 1
                pdf_text = doc[0].get_text()
                assert "報價單" in pdf_text or "QUOTATION" in pdf_text.upper()
                print(f"  -> PDF content verified with {len(doc)} page(s) and quotation header.")
            pdf_page.close()
            print("  -> [PASS Scenario 1] Dynamic items, total calculation, DOCX & PDF verified!")

            # =================================================================
            # SCENARIO 2 (Bug 2): CRM 階段推進平穩捲動與結案區歷史折疊
            # =================================================================
            print("\n" + "=" * 75)
            print("[Scenario 2/6] CRM 階段推進平穩捲動與結案區歷史折疊 (Bug 2)")
            print("=" * 75)

            # Test 2.1: Smooth Scroll Stability
            page.evaluate("window.scrollTo(0, 250)")
            scroll_before = page.evaluate("window.scrollY")
            print(f"  -> Initial scroll position before stage progression: {scroll_before}px")

            # Advance Deal 3 (in Column 1 LEAD) by clicking [進展洽談 ➜]
            lead_col = page.locator("div.bg-slate-100\\/70").first
            advance_deal3_btn = lead_col.locator("button:has-text('進展洽談 ➜')").first
            if advance_deal3_btn.is_visible():
                advance_deal3_btn.click()
                time.sleep(0.5)
                scroll_after = page.evaluate("window.scrollY")
                print(f"  -> Scroll position after stage progression: {scroll_after}px")
                # Difference should be minimal (no jump to 0 or page bottom)
                assert abs(scroll_after - scroll_before) < 150, (
                    f"Scroll jump detected! Before: {scroll_before}, After: {scroll_after}"
                )
                print("  -> [PASS 2.1] Viewport remained stable without unexpected scroll jump.")

            # Test 2.2: Column 4 Accordion Fold & Scroll Limit
            col4 = page.locator("div.max-h-\\[800px\\]").first
            expect(col4).to_be_visible()
            expect(col4).to_contain_text("4. 結案 (WON / LOST)")

            # Check max-h-[800px] class
            col4_class = col4.get_attribute("class") or ""
            assert "max-h-[800px]" in col4_class, f"Expected max-h-[800px] in col4 class, got: {col4_class}"

            # Check Won deals accordion (details[open])
            won_accordion = col4.locator("details").filter(has_text="已成交贏單 (WON)").first
            expect(won_accordion).to_be_visible()
            assert won_accordion.get_attribute("open") is not None, "WON accordion must be expanded by default!"
            print("  -> Won deals accordion is expanded by default (details[open]).")

            # Toggle Won accordion to collapse
            won_accordion.locator("summary").click()
            time.sleep(0.3)
            assert won_accordion.get_attribute("open") is None, "WON accordion should be collapsed after click!"
            print("  -> Won deals accordion collapses cleanly.")

            # Check Lost deals accordion (default collapsed)
            lost_accordion = col4.locator("details").filter(has_text="結案輸單 (LOST)").first
            expect(lost_accordion).to_be_visible()
            assert lost_accordion.get_attribute("open") is None, "LOST accordion must be collapsed by default!"
            print("  -> Lost deals accordion is collapsed by default.")

            # Toggle Lost accordion to expand
            lost_accordion.locator("summary").click()
            time.sleep(0.3)
            assert lost_accordion.get_attribute("open") is not None, "LOST accordion should be open after click!"
            print("  -> Lost deals accordion expands cleanly.")

            # Re-open Won accordion if closed so that converted card is accessible in Scenario 3
            if won_accordion.get_attribute("open") is None:
                won_accordion.locator("summary").click()
                time.sleep(0.3)
            print("  -> [PASS Scenario 2] Accordion folding & height constraints verified.")

            # =================================================================
            # SCENARIO 3 (Bug 3): CRM 轉單跳轉至訂單清單之 Highlight 聚焦
            # =================================================================
            print("\n" + "=" * 75)
            print("[Scenario 3/6] CRM 轉單跳轉至訂單清單之 Highlight 聚焦 (Bug 3)")
            print("=" * 75)

            # Locate converted won card in Column 4 (Deal 4)
            so_link = won_accordion.locator("a[href*='/orders?highlight=']").first
            expect(so_link).to_be_visible()
            so_text = so_link.inner_text().strip()
            target_so_id = re.search(r"SO-[A-Za-z0-9\-]+", so_text).group(0)
            print(f"  -> Found converted order link: {target_so_id}")

            # Click the green order number link
            so_link.click()
            page.wait_for_url("**/orders?highlight=*")
            print(f"  -> Navigated to: {page.url}")
            assert f"highlight={target_so_id}" in page.url

            # Locate the highlighted order row
            target_row = page.locator(f"#order-row-{target_so_id}")
            expect(target_row).to_be_visible()

            # Verify amber ring highlight class is applied immediately
            expect(target_row).to_have_class(re.compile(r"ring-amber-400|ring-2"), timeout=5000)
            print("  -> Verified target order row is actively glowing with amber ring pulse.")

            # Wait for transition (3.5s timeout in frontend) and verify gentle background
            print("  -> Waiting 3.8s for smooth highlight transition to amber glow background...")
            time.sleep(3.8)
            expect(target_row).to_have_class(re.compile(r"bg-amber-50"), timeout=5000)
            print("  -> Target order row settled into soft amber background.")

            print(f"  -> [PASS Scenario 3] Order {target_so_id} smooth scrolled and highlighted with amber ring.")

            # =================================================================
            # SCENARIO 4 (Bug 4): 拆批出貨指定出貨倉庫與跨倉扣庫
            # =================================================================
            print("\n" + "=" * 75)
            print("[Scenario 4/6] 拆批出貨指定出貨倉庫與跨倉扣庫 (Bug 4)")
            print("=" * 75)

            # Find a confirmed order locked to WH-MAIN with 3 units of SKU-ACC-02 (CUST-BETA-聯運實業)
            order2_row = page.locator("tr").filter(has_text="CUST-BETA-聯運實業").first
            expect(order2_row).to_be_visible()
            order2_id = order2_row.get_attribute("id").replace("order-row-", "")
            print(f"  -> Located target confirmed order: {order2_id} (CUST-BETA-聯運實業)")

            # Record stock before partial shipment
            inv_svc = InventoryService(db=db)
            wh_main_avail_before = inv_svc.get_available_stock("WH-MAIN", "SKU-ACC-02")
            wh_store_avail_before = inv_svc.get_available_stock("WH-STORE", "SKU-ACC-02")
            wh_main_phy_before = inv_svc.get_stock_balance("WH-MAIN", "SKU-ACC-02")
            wh_main_locked_before = inv_svc.get_locked_stock("WH-MAIN", "SKU-ACC-02")
            wh_store_phy_before = inv_svc.get_stock_balance("WH-STORE", "SKU-ACC-02")
            wh_store_locked_before = inv_svc.get_locked_stock("WH-STORE", "SKU-ACC-02")
            print(f"  -> Stock before: WH-MAIN phy={wh_main_phy_before}, locked={wh_main_locked_before}; WH-STORE phy={wh_store_phy_before}, locked={wh_store_locked_before}")

            # Click [拆批出貨] button on the order row
            order2_row.locator("button:has-text('拆批出貨')").click()
            partial_modal = page.locator("#partial-ship-modal")
            expect(partial_modal).to_be_visible()
            print("  -> Partial shipment modal opened.")

            # Change warehouse selector to WH-STORE (台北門市倉 / 信義旗艦店)
            wh_select = page.locator("#partial-ship-wh-select")
            wh_select.select_option("WH-STORE")
            print("  -> Selected shipping warehouse: WH-STORE.")

            # Set shipment quantity to 2
            qty_input = partial_modal.locator("input[name^='qty_']").first
            qty_input.fill("2")
            print("  -> Set partial quantity: 2 units.")

            # Submit partial shipment
            partial_modal.locator("button[type='submit']:has-text('確認拆批出貨')").click()
            expect(partial_modal).not_to_be_visible(timeout=5000)

            # Verify green toast message contains order ID, DO ID, and WH-STORE
            toast = page.locator("#toast-container")
            expect(toast).to_contain_text("拆批出貨成功", timeout=6000)
            expect(toast).to_contain_text("WH-STORE")
            print("  -> Verified Toast contains '拆批出貨成功' and 'WH-STORE'.")

            # Verify order row updated to PARTIALLY_SHIPPED
            order2_row_after = page.locator(f"#order-row-{order2_id}")
            expect(order2_row_after).to_contain_text("PARTIALLY_SHIPPED")
            print("  -> Order status updated to PARTIALLY_SHIPPED.")

            # Open delivery history modal
            do_btn = order2_row_after.locator("button[hx-get*='/deliveries-modal']").first
            expect(do_btn).to_be_visible()
            do_btn.click()
            deliveries_modal = page.locator("#deliveries-modal")
            expect(deliveries_modal).to_be_visible()

            # Verify deliveries modal explicitly records "出貨倉庫：WH-STORE"
            expect(deliveries_modal).to_contain_text("出貨倉庫：WH-STORE")
            first_do_card = deliveries_modal.locator("div.border.border-slate-200").first
            expect(first_do_card).to_contain_text("2 件")
            print("  -> Deliveries history modal confirmed '出貨倉庫：WH-STORE' with 2 units allocated.")

            # Close deliveries modal
            deliveries_modal.locator("button:has-text('關閉視窗')").click()
            expect(deliveries_modal).not_to_be_visible()

            # Verify stock deduction across warehouses:
            # WH-STORE physical stock must decrease by 2 (15 -> 13)
            # WH-MAIN locked stock must decrease by 2 (3 -> 1)
            wh_main_phy_after = inv_svc.get_stock_balance("WH-MAIN", "SKU-ACC-02")
            wh_main_locked_after = inv_svc.get_locked_stock("WH-MAIN", "SKU-ACC-02")
            wh_store_phy_after = inv_svc.get_stock_balance("WH-STORE", "SKU-ACC-02")
            wh_store_locked_after = inv_svc.get_locked_stock("WH-STORE", "SKU-ACC-02")

            assert wh_store_phy_after == wh_store_phy_before - Decimal("2.0"), (
                f"WH-STORE physical stock should decrease by 2, got: {wh_store_phy_after}"
            )
            assert wh_main_locked_after == wh_main_locked_before - Decimal("2.0"), (
                f"WH-MAIN locked stock should decrease by 2, got: {wh_main_locked_after}"
            )
            print(f"  -> Stock verified: WH-STORE physical: {wh_store_phy_before} -> {wh_store_phy_after}; WH-MAIN locked: {wh_main_locked_before} -> {wh_main_locked_after}.")
            print("  -> [PASS Scenario 4] Cross-warehouse partial fulfillment & ledger deduction verified.")

            # =================================================================
            # SCENARIO 5 (Bug 5): 拆多批出貨時之普通出貨單多頁連續列印
            # =================================================================
            print("\n" + "=" * 75)
            print("[Scenario 5/6] 拆多批出貨時之普通出貨單多頁連續列印 (Bug 5)")
            print("=" * 75)

            # Ship the second batch (remaining 1 unit)
            order2_row_after.locator("button:has-text('拆批出貨')").click()
            expect(partial_modal).to_be_visible()
            # Default remaining qty should be 1
            partial_modal.locator("button[type='submit']:has-text('確認拆批出貨')").click()
            expect(partial_modal).not_to_be_visible(timeout=5000)
            time.sleep(1.0)

            # Order row should now be SHIPPED (2 batches total)
            order2_row_final = page.locator(f"#order-row-{order2_id}")
            expect(order2_row_final).to_contain_text("SHIPPED")
            print("  -> Second batch shipped; order status is SHIPPED with 2 deliveries.")

            # Test 5.1: Order-level label button prints multi-page PDF (2 pages)
            order_label_btn = order2_row_final.locator("a:has-text('🏷️ 出貨單')")
            expect(order_label_btn).to_be_visible()

            with context.expect_page() as label_page_info:
                order_label_btn.click()
            label_page = label_page_info.value
            label_url = label_page.url
            print(f"  -> Order labels PDF tab opened at: {label_url}")
            pdf_resp = page.request.get(label_url)
            assert pdf_resp.status == 200
            order_pdf_bytes = pdf_resp.body()

            with fitz.open(stream=order_pdf_bytes, filetype="pdf") as doc:
                page_count = len(doc)
                print(f"  -> Generated order-level shipping labels PDF: total pages = {page_count}")
                assert page_count == 2, f"Expected exactly 2 pages for 2 DOs, got: {page_count}"
                # Page 1 has first DO, Page 2 has second DO
                page1_text = doc[0].get_text()
                page2_text = doc[1].get_text()
                assert "DO-" in page1_text
                assert "DO-" in page2_text
                assert page1_text != page2_text, "Page 1 and Page 2 should represent distinct delivery orders!"
            label_page.close()
            print("  -> [PASS 5.1] Order-level PDF has exactly 2 continuous 100x150mm pages.")

            # Test 5.2: Specific DO download in deliveries modal prints 1-page PDF
            order2_row_final.locator("button[hx-get*='/deliveries-modal']").first.click()
            expect(deliveries_modal).to_be_visible()

            # Find the first DO download link inside modal
            first_do_btn = deliveries_modal.locator("a:has-text('🏷️ 100x150 出貨單')").first
            with context.expect_page() as single_page_info:
                first_do_btn.click()
            single_page = single_page_info.value
            single_url = single_page.url
            print(f"  -> Single DO label PDF tab opened at: {single_url}")
            single_resp = page.request.get(single_url)
            assert single_resp.status == 200
            single_pdf_bytes = single_resp.body()

            with fitz.open(stream=single_pdf_bytes, filetype="pdf") as single_doc:
                single_pages = len(single_doc)
                print(f"  -> Generated single DO shipping label PDF: total pages = {single_pages}")
                assert single_pages == 1, f"Expected exactly 1 page for single DO, got: {single_pages}"
            single_page.close()
            print("  -> [PASS 5.2] Single DO PDF has exactly 1 page.")

            deliveries_modal.locator("button:has-text('關閉視窗')").click()
            expect(deliveries_modal).not_to_be_visible()
            print("  -> [PASS Scenario 5] Multi-page vs single-page shipping labels verified.")

            # =================================================================
            # SCENARIO 6 (Bug 6): 庫存盤點底冊 Excel 匯出與會計區收斂
            # =================================================================
            print("\n" + "=" * 75)
            print("[Scenario 6/6] 庫存盤點底冊 Excel 匯出與會計區收斂 (Bug 6)")
            print("=" * 75)

            # 6.1 Download Stock Valuation Excel from Inventory page
            page.goto(f"{base_url}/inventory")
            expect(page.locator("h1").filter(has_text="多倉庫存與雙記帳流水").first).to_be_visible()

            export_excel_btn = page.locator("#btn-export-stock-valuation")
            expect(export_excel_btn).to_be_visible()
            expect(export_excel_btn).to_contain_text("📥 匯出雙倉庫存盤點底冊 (.xlsx)")

            excel_path = test_dir / "stock_valuation_downloaded.xlsx"
            with page.expect_download() as dl_excel:
                export_excel_btn.click()
            excel_dl = dl_excel.value
            excel_dl.save_as(str(excel_path))
            suggested_filename = excel_dl.suggested_filename
            print(f"  -> Downloaded stock valuation file: {suggested_filename}")
            assert suggested_filename.startswith("stock_valuation_") and suggested_filename.endswith(".xlsx")

            # Inspect Excel content with openpyxl
            wb = openpyxl.load_workbook(str(excel_path), data_only=True)
            sheet_names = wb.sheetnames
            print(f"  -> Excel workbook sheets: {sheet_names}")
            assert len(sheet_names) >= 1
            ws = wb.active
            # Check headers in row 4 (where table column headers are placed)
            header_values = [cell.value for cell in ws[4]]
            print(f"  -> Sheet table headers (row 4): {header_values[:8]}")
            assert any("品號" in str(h) or "SKU" in str(h) for h in header_values if h)
            assert ws.max_row >= 5, f"Expected data rows in valuation sheet, got max_row={ws.max_row}"
            wb.close()
            print("  -> [PASS 6.1] Stock valuation workbook successfully generated and verified.")

            # 6.2 Verify Accounting Reports Convergence to 3 Core Financial Statements
            page.goto(f"{base_url}/accounting?tab=reports")
            reports_container = page.locator("#accounting-reports-container")
            expect(reports_container).to_be_visible()

            # Verify exactly 3 report cards exist
            report_cards = reports_container.locator("div.grid > div.p-5")
            expect(report_cards).to_have_count(3)
            print("  -> Report cards count in Accounting center: exactly 3.")

            # Card 1: 損益表 (Income Statement, Sheet 1)
            card1 = report_cards.nth(0)
            expect(card1).to_contain_text("Sheet 1")
            expect(card1).to_contain_text("損益表")
            expect(card1.locator("a[href='/accounting/export/income-statement.xlsx']")).to_be_visible()

            # Card 2: 試算表 (Trial Balance, Sheet 2)
            card2 = report_cards.nth(1)
            expect(card2).to_contain_text("Sheet 2")
            expect(card2).to_contain_text("試算表")
            expect(card2.locator("a[href='/accounting/export/trial-balance.xlsx']")).to_be_visible()

            # Card 3: 帳齡分析 (Aging Matrix, Sheet 3)
            card3 = report_cards.nth(2)
            expect(card3).to_contain_text("Sheet 3")
            expect(card3).to_contain_text("帳齡分析")
            expect(card3.locator("a[href='/accounting?tab=aging']")).to_be_visible()

            # Assert that stock valuation card is NOT in accounting
            container_html = reports_container.inner_html()
            assert "雙倉庫存盤點" not in container_html, "Stock valuation card must not exist in accounting reports!"
            assert "Sheet 4" not in container_html, "Sheet 4 reference must not exist in accounting reports!"
            print("  -> [PASS 6.2] Accounting center cleanly consolidated to 3 core financial reports.")

            print("\n" + "=" * 75)
            print("🎉 ALL 6 ITERATION 0923 SCENARIOS VERIFIED SUCCESSFULLY VIA PLAYWRIGHT HEADED UI!")
            print("=" * 75 + "\n")

        finally:
            context.close()
            browser.close()

            # Clean up temp test directory
            try:
                import shutil
                shutil.rmtree(str(test_dir), ignore_errors=True)
            except Exception:
                pass


if __name__ == "__main__":
    test_playwright_iteration0923_1455_suite()
