from __future__ import annotations
import io
import json
import os
import socket
import sys
import threading
import time
from decimal import Decimal
from pathlib import Path

import docx
import fitz  # PyMuPDF
import uvicorn
from playwright.sync_api import sync_playwright, expect

# Ensure UTF-8 output encoding for Windows PowerShell
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT))

from src.shared.database import Database
from src.shared.event_bus import EventBus
from src.core.domain.order import OrderChannel, OrderStatus
from src.core.services.inventory_service import InventoryService
from src.core.services.order_service import OrderService
from src.core.services.procurement_service import ProcurementService
from src.extended_core.crm.services.crm_service import CrmService
from src.channel_ui.line_bot.flex_templates import (
    build_order_receipt_flex,
    build_order_status_flex,
)
from src.channel_ui.web.app import create_app


def get_free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.bind(("", 0))
        return s.getsockname()[1]


def start_server_in_thread(app, port: int):
    """Start uvicorn server in a background daemon thread."""
    config = uvicorn.Config(app, host="127.0.0.1", port=port, log_level="warning")
    server = uvicorn.Server(config)
    t = threading.Thread(target=server.run, daemon=True)
    t.start()
    return server


def wait_for_server(url: str, timeout: float = 10.0):
    """Poll URL until server responds."""
    import urllib.request
    start = time.time()
    while time.time() - start < timeout:
        try:
            with urllib.request.urlopen(f"{url}/", timeout=1.0) as resp:
                if resp.status == 200:
                    return True
        except Exception:
            time.sleep(0.3)
    raise TimeoutError(f"Server at {url} did not respond within {timeout}s")


def test_playwright_iteration0924_suite():
    """Playwright E2E UI verification for all improvement items in iteration-0924-high-level-name.md.
    MANDATORY HEADED MODE:
      - Batch 1 (Inventory & Dashboard):
          1. Dashboard Low Stock cards: Product Name is prominent bold title; SKU is secondary monospace pill.
          2. Dashboard Recent Orders: Header renamed to '客戶名稱 / 下單時間'.
          3. Inventory Stock Balances table: Column order '商品名稱' precedes '品號 (SKU)'.
      - Batch 2 (Procurement Center):
          1. Procurement creation form: Dropdown options display '商品名稱 (SKU) - 目前成本: NT$...'.
          2. Table header renamed to '採購物料 (品名 / SKU)'.
          3. PO ledger row: Product name in bold precedes muted monospace SKU in parentheses.
          4. Receive PO Modal: Destination warehouse shows '倉庫名 (ID)'; table header '品名 / 品號 (SKU)'.
          5. BOM price comparison results: Table header '物料名稱 / 規格' precedes '料號 (SKU)'.
      - Batch 3 (Sales Orders & Fulfillment Modals):
          1. Orders table: Header renamed to '客戶名稱 / 下單時間'; row displays '客戶名稱 (CUST-xxxx)'.
          2. Ship Confirm Modal: Header displays Customer Name with ID; item table header '商品品名 / SKU'; Product Name bold first.
          3. Partial Ship Modal: Header displays Customer Name with ID; item table header '商品品名 / SKU'; Product Name bold first.
          4. Deliveries Modal: Header displays Customer Name with ID; item table header '商品品名 (SKU)'; allocation shows Product Name.
      - Batch 4 (External Channels, Quotation Documents & Accounting):
          1. Accounting Aging: AR table header '客戶名稱' precedes 'AR 編號'; AP table header '供應商名稱' precedes 'AP 編號'.
          2. Quotation Documents: Downloaded DOCX & PDF tables show '商品名稱' before '規格料號 (SKU)'.
          3. LINE Bot Flex Messages: Receipt & Status bubbles render natural product names before SKU.
    """
    port = get_free_port()
    base_url = f"http://127.0.0.1:{port}"

    test_dir = PROJECT_ROOT / "temp_playwright_0924"
    test_dir.mkdir(parents=True, exist_ok=True)
    test_db_path = test_dir / "test_iteration0924.db"
    bom_temp_file = test_dir / "temp_bom_sample_0924.xlsx"

    if test_db_path.exists():
        try:
            test_db_path.unlink()
        except Exception:
            pass

    db = Database(test_db_path)
    db.init_db()
    event_bus = EventBus()

    # Create app which automatically initializes schema and seeds demo data
    app = create_app(db=db, event_bus=event_bus)

    server = start_server_in_thread(app, port)
    wait_for_server(base_url)
    print(f"\n[*] Playwright Headed Test Server running at {base_url}")

    with sync_playwright() as p:
        # Mandatory Headed Mode with human-observable slow_mo
        browser = p.chromium.launch(headless=False, slow_mo=300)
        context = browser.new_context(viewport={"width": 1440, "height": 900})
        page = context.new_page()
        page.set_default_timeout(2000)
        page.set_default_navigation_timeout(5000)

        page.on("console", lambda msg: print(f"  [BROWSER CONSOLE] {msg.type}: {msg.text}"))

        try:
            # =================================================================
            # SCENARIO 1 (Batch 1): 庫存與營運儀表板高層級名稱翻轉
            # =================================================================
            print("\n" + "=" * 75)
            print("[Scenario 1/4] Batch 1: 庫存與營運儀表板高層級名稱翻轉 (Inventory & Dashboard)")
            print("=" * 75)

            # 1.1 Dashboard Low Stock Cards & Recent Orders
            page.goto(f"{base_url}/")
            expect(page.locator("h1").filter(has_text="營運即時儀表板").first).to_be_visible()

            # Verify Low Stock Card: Product Name primary bold, SKU secondary monospace
            low_stock_box = page.locator("div.bg-rose-50").filter(has_text="低於安全存量商品清單")
            expect(low_stock_box).to_be_visible()

            # "無線溫濕度傳感節點" must be bold title
            low_stock_title = low_stock_box.locator("div.font-bold.text-slate-900").filter(has_text="無線溫濕度傳感節點").first
            expect(low_stock_title).to_be_visible()

            # SKU badge must be secondary pill
            sku_pill = low_stock_box.locator("span.font-mono").filter(has_text="SKU-IOT-03").first
            expect(sku_pill).to_be_visible()
            print("  -> [PASS] Dashboard Low Stock Card: Product Name '無線溫濕度傳感節點' bold title, SKU secondary.")

            # Verify Recent Orders table header: '客戶名稱 / 下單時間'
            recent_orders_table = page.locator("table").filter(has_text="近期訂單動態").first
            # Or header row in the table
            cust_hdr = page.locator("th").filter(has_text="客戶名稱 / 下單時間").first
            expect(cust_hdr).to_be_visible()
            legacy_cust_hdr = page.locator("th").filter(has_text="客戶代號 / 下單時間")
            expect(legacy_cust_hdr).to_have_count(0)
            print("  -> [PASS] Dashboard Recent Orders: Header verified as '客戶名稱 / 下單時間'.")

            # 1.2 Multi-Warehouse Inventory Stock Table Column Order
            page.goto(f"{base_url}/inventory")
            expect(page.locator("h1").filter(has_text="多倉庫存與雙記帳流水").first).to_be_visible()

            stock_table = page.locator("table[data-testid='table-stock-balances']")
            expect(stock_table).to_be_visible()

            # Verify Header: 商品名稱 precedes 品號 (SKU)
            headers = stock_table.locator("thead th").all_inner_texts()
            # Headers list should contain '商品名稱' and '品號 (SKU)'
            pname_col_idx = -1
            sku_col_idx = -1
            for idx, h_text in enumerate(headers):
                if "商品名稱" in h_text:
                    pname_col_idx = idx
                elif "品號 (SKU)" in h_text:
                    sku_col_idx = idx

            assert pname_col_idx != -1, f"Missing '商品名稱' header, found: {headers}"
            assert sku_col_idx != -1, f"Missing '品號 (SKU)' header, found: {headers}"
            assert pname_col_idx < sku_col_idx, (
                f"Expected '商品名稱' (idx {pname_col_idx}) to precede '品號 (SKU)' (idx {sku_col_idx})"
            )

            # Verify First Row Content: Product Name cell precedes SKU cell
            first_row = stock_table.locator("tbody tr").first
            expect(first_row).to_be_visible()
            row_cells = first_row.locator("td").all_inner_texts()
            # Cell 2 is product name, Cell 3 is SKU
            assert "終端機" in row_cells[1] or "擴充底座" in row_cells[1] or "傳感" in row_cells[1], (
                f"Expected cell 1 to be product name, got: {row_cells[1]}"
            )
            assert "SKU-" in row_cells[2], f"Expected cell 2 to be SKU, got: {row_cells[2]}"
            print(f"  -> [PASS] Inventory Stock Table: Column order verified (商品名稱 col {pname_col_idx} < 品號 col {sku_col_idx}).")

            # =================================================================
            # SCENARIO 2 (Batch 2): 採購中心表單與單據明細名稱翻轉
            # =================================================================
            print("\n" + "=" * 75)
            print("[Scenario 2/4] Batch 2: 採購中心表單與單據明細名稱翻轉 (Procurement Center)")
            print("=" * 75)

            page.goto(f"{base_url}/procurement")
            expect(page.locator("h1").filter(has_text="採購與進貨作業中心").first).to_be_visible()

            # 2.1 Verify PO Form Table Header & Option Format
            po_form = page.locator("form[data-testid='form-create-po']")
            expect(po_form).to_be_visible()

            # Table Header must be '採購物料 (品名 / SKU)'
            po_items_hdr = po_form.locator("th").filter(has_text="採購物料 (品名 / SKU)").first
            expect(po_items_hdr).to_be_visible()
            legacy_po_hdr = po_form.locator("th").filter(has_text="採購物料 (SKU / 品名)")
            expect(legacy_po_hdr).to_have_count(0)

            # Dropdown Option must format as '品名 (SKU) - 目前成本: ...'
            sku_select = po_form.locator("select[name='sku']").first
            sku_options_text = sku_select.locator("option").all_inner_texts()
            # Find option for SKU-PRO-01
            pro01_option = next((opt.strip() for opt in sku_options_text if "SKU-PRO-01" in opt), "")
            assert pro01_option.startswith("智慧旗艦終端機 (SKU-PRO-01)"), (
                f"Option text should start with product name, got: '{pro01_option}'"
            )
            assert "目前成本:" in pro01_option, f"Option text should contain cost, got: '{pro01_option}'"
            print(f"  -> [PASS] PO Form Header and Dropdown Option verified: '{pro01_option}'.")

            # 2.2 Create a New PO to verify PO Row Typography
            page.fill("#input-supplier-name", "華碩電腦股份有限公司")
            sku_select.select_option("SKU-PRO-01")
            po_form.locator("input[name='qty']").first.fill("5")
            po_form.locator("input[name='unit_cost']").first.fill("1450.0")

            # Submit PO creation
            po_form.locator("button[data-testid='btn-create-po']").click()
            expect(page.locator("#toast-container")).to_contain_text("成功建立採購單")

            # Locate newly created PO row in #po-tbody
            new_po_row = page.locator("#po-tbody tr").filter(has_text="華碩電腦股份有限公司").first
            expect(new_po_row).to_be_visible()

            # Assert Product Name is bold font-semibold, SKU in muted monospace
            pname_span = new_po_row.locator("span.font-semibold.text-slate-800").filter(has_text="智慧旗艦終端機").first
            expect(pname_span).to_be_visible()
            sku_span = new_po_row.locator("span.font-mono.text-slate-400").filter(has_text="SKU-PRO-01").first
            expect(sku_span).to_be_visible()
            print("  -> [PASS] PO Row item display: Product Name bold precedes muted monospace SKU.")

            # 2.3 Submit PO & Verify Receive Modal
            # Click [確認下訂] to advance status to SUBMITTED
            submit_btn = new_po_row.locator("button:has-text('確認下訂')").first
            submit_btn.click()
            expect(new_po_row).to_contain_text("待到貨", timeout=5000)

            # Click [驗收入庫] to open receive modal
            receive_btn = new_po_row.locator("button:has-text('驗收入庫')").first
            expect(receive_btn).to_be_visible()
            receive_btn.click()

            receive_modal = page.locator("#receive-po-modal")
            expect(receive_modal).to_be_visible()

            # Verify destination warehouse dropdown shows 'Name (ID)'
            wh_option_text = receive_modal.locator("select[name='warehouse_id'] option").all_inner_texts()
            main_wh_opt = next((w for w in wh_option_text if "WH-MAIN" in w), "")
            assert "台北總倉 (WH-MAIN)" in main_wh_opt, f"Expected '台北總倉 (WH-MAIN)', got '{main_wh_opt}'"

            # Verify table header: '品名 / 品號 (SKU)'
            expect(receive_modal.locator("div").filter(has_text="品名 / 品號 (SKU)").first).to_be_visible()

            # Verify item row: Product Name bold title above SKU
            expect(receive_modal.locator("div.font-bold.text-slate-800").filter(has_text="智慧旗艦終端機").first).to_be_visible()
            expect(receive_modal.locator("div.text-slate-400.font-mono").filter(has_text="SKU-PRO-01").first).to_be_visible()

            # Cancel receive modal
            receive_modal.locator("button:has-text('取消')").first.click()
            expect(receive_modal).not_to_be_visible()
            print("  -> [PASS] Receive PO Modal: Warehouse Name (ID), '品名 / 品號 (SKU)' header, and Product Name verified.")

            # 2.4 Verify BOM Price Comparison Table Header
            page.locator("#tab-btn-bom").click()
            bom_section = page.locator("#section-bom-po")
            expect(bom_section).to_be_visible()

            # Download sample BOM Excel
            with page.expect_download() as dl_bom_info:
                bom_section.locator("a[href='/procurement/bom/sample']").click()
            bom_download = dl_bom_info.value
            bom_download.save_as(str(bom_temp_file))

            # Upload BOM Excel file and compare
            bom_file_input = bom_section.locator("input[name='file']")
            bom_file_input.set_input_files(str(bom_temp_file))
            bom_section.locator("button:has-text('開始智慧比價')").click()

            # Wait for BOM comparison result table
            split_btn = page.locator("#btn-split-bom-pos")
            expect(split_btn).to_be_visible(timeout=8000)

            # Verify table headers: '物料名稱 / 規格' precedes '料號 (SKU)'
            bom_headers = page.locator("#section-bom-po table thead th").all_inner_texts()
            b_name_idx = -1
            b_sku_idx = -1
            for idx, h_text in enumerate(bom_headers):
                if "物料名稱" in h_text:
                    b_name_idx = idx
                elif "料號 (SKU)" in h_text:
                    b_sku_idx = idx

            assert b_name_idx != -1 and b_sku_idx != -1, f"BOM headers missing expected keys: {bom_headers}"
            assert b_name_idx < b_sku_idx, (
                f"Expected '物料名稱 / 規格' (idx {b_name_idx}) to precede '料號 (SKU)' (idx {b_sku_idx})"
            )
            print(f"  -> [PASS] BOM Comparison Table: Header '物料名稱 / 規格' (col {b_name_idx}) precedes '料號 (SKU)' (col {b_sku_idx}).")

            # =================================================================
            # SCENARIO 3 (Batch 3): 銷售訂單與履約彈窗名稱翻轉
            # =================================================================
            print("\n" + "=" * 75)
            print("[Scenario 3/4] Batch 3: 銷售訂單與履約彈窗名稱翻轉 (Sales Orders & Fulfillment Modals)")
            print("=" * 75)

            page.goto(f"{base_url}/orders")
            expect(page.locator("h1").filter(has_text="訂單與出貨作業中心").first).to_be_visible()

            # 3.1 Orders Table Header & Row Customer Display
            orders_table = page.locator("table[data-testid='table-orders']")
            expect(orders_table).to_be_visible()

            # Verify header renamed
            cust_hdr = orders_table.locator("thead th").filter(has_text="客戶名稱 / 下單時間")
            expect(cust_hdr).to_be_visible()
            expect(orders_table.locator("thead th").filter(has_text="客戶代號 / 下單時間")).to_have_count(0)

            # Verify row customer format: '客戶名稱 (CUST-xxxx)'
            target_order_row = orders_table.locator("tbody tr").filter(has_text="宏達數位科技").first
            expect(target_order_row).to_be_visible()
            target_order_id = target_order_row.locator("td").first.locator("div").first.inner_text().strip()
            target_order_text = target_order_row.inner_text()
            assert "宏達數位科技" in target_order_text
            assert "CUST-" in target_order_text
            print(f"  -> [PASS] Orders Table: Header '客戶名稱 / 下單時間' and Customer Name + ID row verified for {target_order_id}.")

            # 3.2 Ship Confirm Modal Verification
            ship_btn = target_order_row.locator("button:has-text('一鍵出貨 (扣庫)')")
            expect(ship_btn).to_be_visible()
            ship_btn.click()

            ship_modal = page.locator("#ship-confirm-modal")
            expect(ship_modal).to_be_visible()

            # Verify Header Customer display: '客戶：客戶名稱 (ID)'
            modal_header_text = ship_modal.locator("div.p-5").first.inner_text()
            assert "客戶：" in modal_header_text
            assert "CUST-" in modal_header_text

            # Verify Table Header: '商品品名 / SKU'
            expect(ship_modal.locator("div").filter(has_text="商品品名 / SKU").first).to_be_visible()

            # Verify Item rows: Product Name bold precedes muted SKU in parentheses
            item_row = ship_modal.locator(".border.border-slate-200 .grid-cols-12").nth(1)
            pname_div = item_row.locator("div.font-semibold.text-slate-800")
            sku_div = item_row.locator("div.font-mono.text-xs.text-slate-400")
            expect(pname_div).to_be_visible()
            expect(sku_div).to_be_visible()
            assert "(" in sku_div.inner_text() and ")" in sku_div.inner_text()

            # Close ship confirm modal
            ship_modal.locator("button:has-text('取消')").click()
            expect(ship_modal).not_to_be_visible()
            print("  -> [PASS] Ship Confirm Modal: Customer Name (ID), '商品品名 / SKU', and Product Name bold verified.")

            # 3.3 Partial Ship Modal Verification
            partial_btn = target_order_row.locator("button:has-text('拆批出貨')")
            expect(partial_btn).to_be_visible()
            partial_btn.click()

            partial_modal = page.locator("#partial-ship-modal")
            expect(partial_modal).to_be_visible()

            # Verify Header Customer display: '客戶：客戶名稱 (ID)'
            p_modal_header = partial_modal.locator("div.p-5").first.inner_text()
            assert "客戶：" in p_modal_header
            assert "CUST-" in p_modal_header

            # Verify Table Header: '商品品名 / SKU'
            expect(partial_modal.locator("div").filter(has_text="商品品名 / SKU").first).to_be_visible()

            # Verify Item rows: Product Name bold precedes muted SKU in parentheses
            p_item_row = partial_modal.locator(".border.border-slate-200 .grid-cols-12").nth(1)
            expect(p_item_row.locator("div.font-semibold.text-slate-800")).to_be_visible()
            expect(p_item_row.locator("div.font-mono.text-xs.text-slate-400")).to_be_visible()

            # Execute partial shipment of 1 unit to verify deliveries modal next
            qty_input = partial_modal.locator("input[name^='qty_']").first
            qty_input.fill("1")
            partial_modal.locator("button:has-text('確認拆批出貨')").click()
            expect(page.locator("#toast-container")).to_contain_text("拆批出貨成功")
            print("  -> [PASS] Partial Ship Modal verified and executed partial shipment.")

            # 3.4 Deliveries History Modal Verification
            # Order row should now show delivery button
            updated_order_row = orders_table.locator(f"#order-row-{target_order_id}")
            deliv_btn = updated_order_row.locator("button[hx-get*='deliveries-modal']").first
            expect(deliv_btn).to_be_visible()
            deliv_btn.click()

            deliv_modal = page.locator("#deliveries-modal")
            expect(deliv_modal).to_be_visible()

            # Verify Header Customer display
            deliv_hdr_text = deliv_modal.locator("div.p-5").first.inner_text()
            assert "客戶：" in deliv_hdr_text
            assert "宏達數位科技" in deliv_hdr_text
            assert "CUST-" in deliv_hdr_text

            # Verify Allocation Table Header: '商品品名 (SKU)'
            expect(deliv_modal.locator("th").filter(has_text="商品品名 (SKU)")).to_be_visible()

            # Verify Allocation Item Row: Product Name bold precedes (SKU)
            alloc_item = deliv_modal.locator("tbody tr").first
            expect(alloc_item.locator("span.font-semibold.text-slate-900")).to_be_visible()
            expect(alloc_item.locator("span.font-mono")).to_be_visible()

            # Close deliveries modal
            deliv_modal.locator("button:has-text('關閉視窗')").click()
            expect(deliv_modal).not_to_be_visible()
            print("  -> [PASS] Deliveries Modal: Customer Name (ID), '商品品名 (SKU)' header, and Allocation Product Name verified.")

            # =================================================================
            # SCENARIO 4 (Batch 4): 外部通路、報價單與會計名稱翻轉
            # =================================================================
            print("\n" + "=" * 75)
            print("[Scenario 4/4] Batch 4: 外部通路、報價單與會計名稱翻轉 (LINE, CRM Documents & Accounting)")
            print("=" * 75)

            # 4.1 Accounting Aging Tables: Customer/Supplier Name before ID
            page.goto(f"{base_url}/accounting?tab=aging")
            expect(page.locator("h1").filter(has_text="探吉 會計與記帳系統").first).to_be_visible()

            # AR Aging Table: 客戶名稱 header must precede AR 編號
            ar_table = page.locator("table").filter(has_text="AR 編號").first
            expect(ar_table).to_be_visible()
            ar_headers = ar_table.locator("thead th").all_inner_texts()
            ar_cust_idx = -1
            ar_id_idx = -1
            for idx, h in enumerate(ar_headers):
                if "客戶名稱" in h:
                    ar_cust_idx = idx
                elif "AR 編號" in h:
                    ar_id_idx = idx
            assert ar_cust_idx != -1 and ar_id_idx != -1, f"AR table missing expected headers: {ar_headers}"
            assert ar_cust_idx < ar_id_idx, (
                f"AR table: '客戶名稱' (idx {ar_cust_idx}) must precede 'AR 編號' (idx {ar_id_idx})"
            )

            # AP Aging Table: 供應商名稱 header must precede AP 編號
            ap_table = page.locator("table").filter(has_text="AP 編號").first
            expect(ap_table).to_be_visible()
            ap_headers = ap_table.locator("thead th").all_inner_texts()
            ap_supp_idx = -1
            ap_id_idx = -1
            for idx, h in enumerate(ap_headers):
                if "供應商名稱" in h:
                    ap_supp_idx = idx
                elif "AP 編號" in h:
                    ap_id_idx = idx
            assert ap_supp_idx != -1 and ap_id_idx != -1, f"AP table missing expected headers: {ap_headers}"
            assert ap_supp_idx < ap_id_idx, (
                f"AP table: '供應商名稱' (idx {ap_supp_idx}) must precede 'AP 編號' (idx {ap_id_idx})"
            )
            print(f"  -> [PASS] Accounting Aging: AR table (客戶名稱 col {ar_cust_idx} < AR編號 col {ar_id_idx}) and AP table verified.")

            # 4.2 CRM Quotation Documents (DOCX & PDF): Item Table Header & Column Order
            page.goto(f"{base_url}/crm?tab=kanban")
            expect(page.locator("h1").filter(has_text="哩來 (LiLai) CRM").first).to_be_visible()

            # Locate deal card with quotation (Deal 1 or Deal 4)
            quote_card = page.locator("div[id^='deal-card-']").filter(has_text="商業報價單").first
            expect(quote_card).to_be_visible()

            # Download DOCX
            docx_temp_path = test_dir / "downloaded_quote.docx"
            with page.expect_download() as dl_docx_info:
                quote_card.locator("a:has-text('📥 DOCX')").click()
            dl_docx = dl_docx_info.value
            dl_docx.save_as(str(docx_temp_path))
            assert docx_temp_path.exists()

            # Verify DOCX table headers: ["項次", "商品名稱", "規格料號 (SKU)", "數量", "單價 (NT$)", "小計 (NT$)"]
            doc = docx.Document(str(docx_temp_path))
            item_table = next((t for t in doc.tables if len(t.columns) == 6), None)
            assert item_table is not None, "Could not find 6-column line items table in generated DOCX"
            docx_headers = [item_table.cell(0, i).text.strip() for i in range(6)]
            assert docx_headers[1] == "商品名稱", f"Expected '商品名稱' at col 1, got '{docx_headers[1]}'"
            assert docx_headers[2] == "規格料號 (SKU)", f"Expected '規格料號 (SKU)' at col 2, got '{docx_headers[2]}'"

            # Preview PDF & inspect text (target='_blank' inline PDF opens new tab)
            pdf_link = quote_card.locator("a:has-text('📑 預覽 PDF')")
            with context.expect_page() as new_page_info:
                pdf_link.click()
            pdf_page = new_page_info.value
            pdf_url = pdf_page.url
            pdf_resp = page.request.get(pdf_url)
            assert pdf_resp.status == 200
            pdf_bytes = pdf_resp.body()
            pdf_doc = fitz.open(stream=pdf_bytes, filetype="pdf")
            pdf_text = pdf_doc[0].get_text()
            p_name_pos = pdf_text.find("商品名稱")
            p_sku_pos = pdf_text.find("規格料號 (SKU)")
            assert p_name_pos != -1 and p_sku_pos != -1
            assert p_name_pos < p_sku_pos, (
                f"PDF text: '商品名稱' (pos {p_name_pos}) must precede '規格料號 (SKU)' (pos {p_sku_pos})"
            )
            pdf_page.close()
            print("  -> [PASS] CRM Quote Documents: Both DOCX and PDF verified with '商品名稱' preceding '規格料號 (SKU)'.")

            # 4.3 LINE Bot Flex Message Product Name Priority
            line_items = [
                {"sku": "SKU-PRO-01", "product_name": "智慧旗艦終端機", "qty": 2, "subtotal": 5000.0}
            ]
            receipt_flex = build_order_receipt_flex(
                order_id="SO-LINE-VERIFY",
                total_amount=Decimal("5000.0"),
                items=line_items,
            )
            receipt_json = json.dumps(receipt_flex.contents.to_dict(), ensure_ascii=False)
            assert "智慧旗艦終端機 (SKU-PRO-01) x 2" in receipt_json, (
                f"Receipt flex missing product name priority: {receipt_json}"
            )

            status_flex = build_order_status_flex(
                order_id="SO-LINE-VERIFY",
                status="CONFIRMED",
                total_amount=5000.0,
                items=line_items,
            )
            status_json = json.dumps(status_flex.contents.to_dict(), ensure_ascii=False)
            assert "智慧旗艦終端機 (SKU-PRO-01) x 2" in status_json, (
                f"Status flex missing product name priority: {status_json}"
            )
            print("  -> [PASS] LINE Bot Flex Messages: Receipt & Status bubbles verified with '品名 (SKU) x 數量'.")

            print("\n" + "=" * 75)
            print("🎉 ALL 4 SCENARIOS IN ITERATION 0924 VERIFIED SUCCESSFULLY!")
            print("=" * 75 + "\n")

        finally:
            browser.close()
            server.should_exit = True
            time.sleep(0.5)
            for f in test_dir.glob("*"):
                try:
                    f.unlink()
                except Exception:
                    pass
            try:
                test_dir.rmdir()
            except Exception:
                pass


if __name__ == "__main__":
    test_playwright_iteration0924_suite()
