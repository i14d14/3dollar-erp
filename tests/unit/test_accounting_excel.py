import datetime
import io
import unittest
from decimal import Decimal
import openpyxl

from src.extended_core.accounting.models import (
    AccountType,
    IncomeStatement,
    TrialBalance,
    TrialBalanceItem,
)
from src.extended_core.accounting.reports.excel_exporter import (
    export_income_statement_xlsx,
    export_trial_balance_xlsx,
)


class TestAccountingExcelExport(unittest.TestCase):
    def test_export_trial_balance_xlsx(self):
        items = [
            TrialBalanceItem(
                account_code="1101",
                account_name="現金",
                account_type=AccountType.ASSET,
                debit_total=Decimal("10000.0"),
                credit_total=Decimal("0.0"),
                balance_debit=Decimal("10000.0"),
                balance_credit=Decimal("0.0"),
            ),
            TrialBalanceItem(
                account_code="3101",
                account_name="業主資本",
                account_type=AccountType.EQUITY,
                debit_total=Decimal("0.0"),
                credit_total=Decimal("10000.0"),
                balance_debit=Decimal("0.0"),
                balance_credit=Decimal("10000.0"),
            ),
        ]
        tb = TrialBalance(
            as_of_date=datetime.date.today(),
            items=items,
            total_debit=Decimal("10000.0"),
            total_credit=Decimal("10000.0"),
            is_balanced=True,
        )

        xlsx_bytes = export_trial_balance_xlsx(tb)
        self.assertGreater(len(xlsx_bytes), 1000)

        # Parse with openpyxl to verify worksheet structure
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        self.assertIn("會計試算表", wb.sheetnames)
        ws = wb["會計試算表"]
        # Check title in row 1
        self.assertEqual(ws["A1"].value, "三圓科技股份有限公司")
        # Check column headers
        self.assertEqual(ws["A5"].value, "科目代碼")
        # Check first item
        self.assertEqual(ws["A6"].value, "1101")
        self.assertEqual(ws["B6"].value, "現金")

    def test_export_income_statement_xlsx(self):
        stmt = IncomeStatement(
            end_date=datetime.date.today(),
            revenue=Decimal("100000.0"),
            cogs=Decimal("60000.0"),
            gross_profit=Decimal("40000.0"),
            operating_expenses=Decimal("15000.0"),
            net_profit=Decimal("25000.0"),
        )
        xlsx_bytes = export_income_statement_xlsx(stmt)
        self.assertGreater(len(xlsx_bytes), 1000)

        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        self.assertIn("損益表", wb.sheetnames)
        ws = wb["損益表"]
        self.assertEqual(ws["A1"].value, "三圓科技股份有限公司")
        self.assertEqual(ws["A6"].value, "營業收入 (Sales Revenue)")
        self.assertEqual(ws["B6"].value, 100000.0)


if __name__ == "__main__":
    unittest.main()
