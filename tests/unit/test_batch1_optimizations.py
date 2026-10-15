from pathlib import Path
from decimal import Decimal
import io
import openpyxl
import pytest
from starlette.testclient import TestClient

from src.channel_ui.web.app import create_app
from src.core.domain.inventory import Product
from src.core.services.inventory_service import InventoryService
from src.extended_core.crm.models import DealStage
from src.extended_core.crm.services.crm_service import CrmService
from src.shared.database import Database
from src.shared.event_bus import EventBus


@pytest.fixture
def client_and_services(tmp_path):
    """Create test database, services, and web test client."""
    db_file = tmp_path / "test_batch1.db"
    db = Database(db_file)
    event_bus = EventBus()
    app = create_app(db=db, event_bus=event_bus)
    client = TestClient(app, raise_server_exceptions=False)

    inv_service = InventoryService(db=db)
    crm_service = CrmService(db=db, event_bus=event_bus)

    # Add extra test product
    inv_service.create_product(sku="SKU-C", name="物料 C", base_unit="PCS", cost_price=Decimal("300.00"), selling_price=Decimal("600.00"))

    yield client, db, crm_service, inv_service

    if db_file.exists():
        try:
            db_file.unlink()
        except Exception:
            pass


def test_quote_dynamic_multi_items_creation(client_and_services):
    """Bug 1: Verify quotation creation supports 3+ dynamic items."""
    client, db, crm_service, _ = client_and_services

    contact = crm_service.create_contact(company_name="測試企業", contact_name="王小明")
    deal = crm_service.create_deal(contact_id=contact.contact_id, title="多品項採購案", stage=DealStage.DISCOVERY)

    # Post with 3 dynamic items using form lists: sku, qty, price
    payload = {
        "sku": ["SKU-PRO-01", "SKU-ACC-02", "SKU-IOT-03"],
        "qty": ["2", "3", "1"],
        "price": ["2500.00", "650.00", "800.00"],
        "valid_days": "45",
    }
    resp = client.post(f"/crm/deals/{deal.deal_id}/quotes", data=payload, headers={"HX-Request": "true"})
    assert resp.status_code == 200

    # Verify deal advanced to PROPOSAL
    updated_deal = crm_service.get_deal(deal.deal_id)
    assert updated_deal.stage == DealStage.PROPOSAL

    # Verify quotation in DB has 3 items
    quotes = crm_service.get_quotes_for_deal(deal.deal_id)
    assert len(quotes) == 1
    quote = quotes[0]
    assert len(quote.items) == 3
    # Expected total: 2*2500 + 3*650 + 1*800 = 5000 + 1950 + 800 = 7750
    assert quote.total_amount == Decimal("7750.00")
    assert quote.valid_until is not None


def test_quote_legacy_fallback_compatibility(client_and_services):
    """Bug 1: Verify fallback for legacy form submitting sku_1 and sku_2."""
    client, db, crm_service, _ = client_and_services

    contact = crm_service.create_contact(company_name="相容性企業", contact_name="李大華")
    deal = crm_service.create_deal(contact_id=contact.contact_id, title="舊版相容性商機", stage=DealStage.DISCOVERY)

    payload = {
        "sku_1": "SKU-PRO-01",
        "qty_1": "5",
        "price_1": "2000.00",
        "sku_2": "SKU-ACC-02",
        "qty_2": "2",
        "price_2": "600.00",
        "valid_days": "30",
    }
    resp = client.post(f"/crm/deals/{deal.deal_id}/quotes", data=payload, headers={"HX-Request": "true"})
    assert resp.status_code == 200

    quotes = crm_service.get_quotes_for_deal(deal.deal_id)
    assert len(quotes) == 1
    quote = quotes[0]
    assert len(quote.items) == 2
    # 5*2000 + 2*600 = 10000 + 1200 = 11200
    assert quote.total_amount == Decimal("11200.00")


def test_inventory_stock_valuation_excel_export(client_and_services):
    """Bug 6: Verify standalone stock inventory valuation excel export from /inventory."""
    client, db, _, inv_service = client_and_services

    # Add some stock
    inv_service.adjust_stock("WH-MAIN", "SKU-PRO-01", Decimal("50"), Decimal("1500"), "INIT-01")
    inv_service.adjust_stock("WH-STORE", "SKU-ACC-02", Decimal("20"), Decimal("300"), "INIT-02")

    resp = client.get("/inventory/export/stock-valuation.xlsx")
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers["content-type"]
    assert "stock_valuation_" in resp.headers["content-disposition"]

    # Parse with openpyxl to verify content structure
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert "庫存盤點底冊 (Stock Valuation)" in wb.sheetnames
    ws = wb["庫存盤點底冊 (Stock Valuation)"]

    # Title & headers
    assert "3Dollar ERP - 雙倉庫存盤點與估值底冊" in ws.cell(row=1, column=1).value
    headers = [ws.cell(row=4, column=col).value for col in range(1, 9)]
    expected_headers = ["倉庫代碼", "倉庫名稱", "品號 (SKU)", "品名", "計量單位", "庫存數量", "移動平均單價", "庫存總估值"]
    assert headers == expected_headers


def test_ui_templates_contain_required_features(client_and_services):
    """Verify Bug 1, 2, 3, 6 markup in rendered templates."""
    client, _, _, _ = client_and_services

    # 1. CRM template check
    crm_resp = client.get("/crm?tab=kanban")
    assert crm_resp.status_code == 200
    html = crm_resp.text
    # Bug 1: Dynamic items template and add button
    assert 'id="quote-items-container"' in html
    assert "＋ 新增報價項目" in html
    assert 'id="quote-item-template"' in html
    # Bug 2: Collapsible Won/Lost section
    assert "已成交贏單 (WON)" in html
    assert "結案輸單 (LOST)" in html
    assert "max-h-[800px]" in html
    # Bug 3: Converted order link with highlight parameter
    assert "orders?highlight=" in html or "id=\"deal-card-" in html

    # 2. Inventory template check
    inv_resp = client.get("/inventory")
    assert inv_resp.status_code == 200
    assert 'href="/inventory/export/stock-valuation.xlsx"' in inv_resp.text
    assert "匯出雙倉庫存盤點底冊 (.xlsx)" in inv_resp.text

    # 3. Orders template check
    orders_resp = client.get("/orders")
    assert orders_resp.status_code == 200
    assert "order-row-${highlightId}" in orders_resp.text
    assert "ring-amber-400" in orders_resp.text

    # 4. Accounting template check (Sheet 4 card removed)
    acc_resp = client.get("/accounting?tab=reports")
    assert acc_resp.status_code == 200
    assert "grid grid-cols-1 md:grid-cols-3" in acc_resp.text
    assert "即時下載財務總帳整合報表" in acc_resp.text
