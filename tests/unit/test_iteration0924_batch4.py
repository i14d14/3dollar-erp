import datetime
from decimal import Decimal
import json
from pathlib import Path
import unittest
import docx
import fitz

from src.channel_ui.line_bot.flex_templates import (
    build_order_receipt_flex,
    build_order_status_flex,
)
from src.extended_core.crm.documents.quote_document_engine import QuoteDocumentEngine
from src.extended_core.crm.models import Contact, Deal, DealStage, Quote, QuoteStatus


class TestIteration0924Batch4(unittest.TestCase):
    """Unit tests for Iteration 0924 Batch 4: High-level name priority in LINE, Documents & Accounting."""

    def setUp(self):
        self.templates_dir = Path(__file__).resolve().parents[2] / "src" / "channel_ui" / "web" / "templates"

    def test_line_flex_order_receipt_and_status_prioritize_product_name(self):
        items_with_name = [
            {"sku": "SKU-TEA-01", "product_name": "阿里山高山烏龍茶", "qty": 2, "subtotal": 1600.0},
        ]
        items_without_name = [
            {"sku": "SKU-GENERIC-02", "qty": 1, "subtotal": 500.0},
        ]

        # 1. Receipt Flex Message with product name
        receipt_flex = build_order_receipt_flex(
            order_id="SO-LINE-001",
            total_amount=Decimal("1600.0"),
            items=items_with_name,
        )
        receipt_json = json.dumps(receipt_flex.contents.to_dict(), ensure_ascii=False)
        self.assertIn("阿里山高山烏龍茶 (SKU-TEA-01) x 2", receipt_json)

        # 2. Receipt Flex Message backward compatibility without product name
        receipt_legacy = build_order_receipt_flex(
            order_id="SO-LINE-002",
            total_amount=Decimal("500.0"),
            items=items_without_name,
        )
        receipt_legacy_json = json.dumps(receipt_legacy.contents.to_dict(), ensure_ascii=False)
        self.assertIn("SKU-GENERIC-02 x 1", receipt_legacy_json)

        # 3. Status Flex Message with product name
        status_flex = build_order_status_flex(
            order_id="SO-LINE-001",
            status="CONFIRMED",
            total_amount=1600.0,
            items=items_with_name,
        )
        status_json = json.dumps(status_flex.contents.to_dict(), ensure_ascii=False)
        self.assertIn("阿里山高山烏龍茶 (SKU-TEA-01) x 2", status_json)

    def test_quote_docx_line_items_table_order(self):
        contact = Contact(
            contact_id="CUST-DOC-001",
            company_name="台積電股份有限公司",
            tax_id="12345678",
            contact_name="張處長",
        )
        deal = Deal(
            deal_id="DEAL-DOC-001",
            contact_id=contact.contact_id,
            title="年度企業禮品採購案",
            expected_amount=Decimal("50000.0"),
            stage=DealStage.PROPOSAL,
        )
        quote = Quote(
            quote_id="Q-DOC-001",
            deal_id=deal.deal_id,
            quote_number="QUO-20260924-001",
            valid_until=datetime.date(2026, 12, 31),
            total_amount=Decimal("50000.0"),
            status=QuoteStatus.SENT,
        )
        items = [
            {"sku": "SKU-GIFT-A", "name": "特級高山茶禮盒", "qty": 50, "unit_price": 1000.0, "subtotal": 50000.0}
        ]

        docx_bytes = QuoteDocumentEngine.generate_docx(contact, deal, quote, items)
        import io
        doc = docx.Document(io.BytesIO(docx_bytes))

        # Find the line items table (second table, with 6 columns)
        item_table = None
        for t in doc.tables:
            if len(t.columns) == 6:
                item_table = t
                break
        self.assertIsNotNone(item_table)

        # Header: ["項次", "商品名稱", "規格料號 (SKU)", "數量", "單價 (NT$)", "小計 (NT$)"]
        headers = [item_table.cell(0, i).text.strip() for i in range(6)]
        self.assertEqual(headers[0], "項次")
        self.assertEqual(headers[1], "商品名稱")
        self.assertEqual(headers[2], "規格料號 (SKU)")

        # Row 1 data: name before sku
        row1_cells = [item_table.cell(1, i).text.strip() for i in range(6)]
        self.assertEqual(row1_cells[1], "特級高山茶禮盒")
        self.assertEqual(row1_cells[2], "SKU-GIFT-A")

    def test_quote_pdf_line_items_table_order(self):
        contact = Contact(
            contact_id="CUST-DOC-001",
            company_name="台積電股份有限公司",
            tax_id="12345678",
            contact_name="張處長",
        )
        deal = Deal(
            deal_id="DEAL-DOC-001",
            contact_id=contact.contact_id,
            title="年度企業禮品採購案",
            expected_amount=Decimal("50000.0"),
            stage=DealStage.PROPOSAL,
        )
        quote = Quote(
            quote_id="Q-DOC-001",
            deal_id=deal.deal_id,
            quote_number="QUO-20260924-001",
            valid_until=datetime.date(2026, 12, 31),
            total_amount=Decimal("50000.0"),
            status=QuoteStatus.SENT,
        )
        items = [
            {"sku": "SKU-GIFT-A", "name": "特級高山茶禮盒", "qty": 50, "unit_price": 1000.0, "subtotal": 50000.0}
        ]

        pdf_bytes = QuoteDocumentEngine.generate_pdf(contact, deal, quote, items)
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        text = doc[0].get_text()

        pname_idx = text.find("商品名稱")
        sku_idx = text.find("規格料號 (SKU)")
        self.assertNotEqual(pname_idx, -1)
        self.assertNotEqual(sku_idx, -1)
        self.assertLess(pname_idx, sku_idx, "商品名稱 header must appear before 規格料號 (SKU) in PDF")

    def test_accounting_aging_tables_prioritize_partner_name(self):
        accounting_path = self.templates_dir / "accounting.html"
        self.assertTrue(accounting_path.exists())
        content = accounting_path.read_text(encoding="utf-8")

        # In AR aging table, 客戶名稱 must precede AR 編號
        ar_table_idx = content.find("{{ ar_aging.report_type }} 帳齡分析與款項催收")
        self.assertNotEqual(ar_table_idx, -1)
        ar_cust_idx = content.find('<th class="py-2.5 px-4">客戶名稱</th>', ar_table_idx)
        ar_id_idx = content.find('<th class="py-2.5 px-4">AR 編號</th>', ar_table_idx)
        self.assertNotEqual(ar_cust_idx, -1)
        self.assertNotEqual(ar_id_idx, -1)
        self.assertLess(ar_cust_idx, ar_id_idx, "客戶名稱 must appear before AR 編號 in AR aging table")

        # In AP aging table, 供應商名稱 must precede AP 編號
        ap_table_idx = content.find("{{ ap_aging.report_type }} 帳齡分析與應付管理")
        self.assertNotEqual(ap_table_idx, -1)
        ap_supp_idx = content.find('<th class="py-2.5 px-4">供應商名稱</th>', ap_table_idx)
        ap_id_idx = content.find('<th class="py-2.5 px-4">AP 編號</th>', ap_table_idx)
        self.assertNotEqual(ap_supp_idx, -1)
        self.assertNotEqual(ap_id_idx, -1)
        self.assertLess(ap_supp_idx, ap_id_idx, "供應商名稱 must appear before AP 編號 in AP aging table")


if __name__ == "__main__":
    unittest.main()
