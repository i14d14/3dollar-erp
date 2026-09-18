import datetime
from decimal import Decimal
import io
import os
from pathlib import Path
import unittest
import fitz
import docx

from src.core.domain.order import OrderChannel, OrderStatus
from src.core.services.inventory_service import InventoryService
from src.core.services.order_service import OrderService
from src.extended_core.crm.models import DealStage, QuoteStatus
from src.extended_core.crm.services.crm_service import CrmService
from src.shared.database import Database
from src.shared.event_bus import EventBus


class TestCrmService(unittest.TestCase):
    def setUp(self):
        self.test_db_path = Path("d:/3dollar_erp/test_crm.db")
        if self.test_db_path.exists():
            self.test_db_path.unlink()
        self.db = Database(self.test_db_path)
        self.db.init_db()

        self.event_bus = EventBus()
        self.inventory_service = InventoryService(db=self.db)
        self.order_service = OrderService(db=self.db, event_bus=self.event_bus, inventory_service=self.inventory_service)
        self.crm_service = CrmService(db=self.db, order_service=self.order_service)

        # Create basic warehouse and product
        self.inventory_service.create_warehouse("WH-01", "測試倉")
        self.inventory_service.create_product(
            sku="SKU-CRM-01",
            name="商務整合伺服器",
            base_unit="台",
            cost_price=Decimal("12000.0"),
            selling_price=Decimal("20000.0"),
            safety_stock=5,
        )
        self.inventory_service.create_product(
            sku="SKU-CRM-02",
            name="雲端連網模組",
            base_unit="組",
            cost_price=Decimal("1500.0"),
            selling_price=Decimal("3500.0"),
            safety_stock=10,
        )

    def tearDown(self):
        if self.test_db_path.exists():
            self.test_db_path.unlink()

    def test_contact_crud_and_tax_id_validation(self):
        # 1. Valid contact with 8-digit tax ID
        contact = self.crm_service.create_contact(
            company_name="三圓數位系統股份有限公司",
            contact_name="林總經理",
            tax_id="24567890",
            email="service@sanyuan.tw",
            phone="02-23456789",
        )
        self.assertTrue(contact.contact_id.startswith("CUST-"))
        self.assertEqual(contact.tax_id, "24567890")

        # 2. Invalid tax ID (less than 8 digits or non-digits)
        with self.assertRaises(ValueError):
            self.crm_service.create_contact(
                company_name="測試企業",
                contact_name="王小明",
                tax_id="12345",  # Invalid
            )

        # 3. Retrieve and list
        retrieved = self.crm_service.get_contact(contact.contact_id)
        self.assertIsNotNone(retrieved)
        self.assertEqual(retrieved.company_name, "三圓數位系統股份有限公司")

        contacts = self.crm_service.list_contacts()
        self.assertEqual(len(contacts), 1)

        # 4. Prevent duplicate company name
        with self.assertRaises(ValueError):
            self.crm_service.create_contact(
                company_name="三圓數位系統股份有限公司",
                contact_name="李經理",
            )

        # 5. Update contact (valid update)
        updated = self.crm_service.update_contact(
            contact_id=contact.contact_id,
            contact_name="林副總裁",
            phone="0912-345-678",
        )
        self.assertEqual(updated.contact_name, "林副總裁")
        self.assertEqual(updated.phone, "0912-345-678")
        self.assertEqual(updated.company_name, "三圓數位系統股份有限公司")

        # 6. Prohibit modifying company name
        with self.assertRaises(ValueError):
            self.crm_service.update_contact(
                contact_id=contact.contact_id,
                company_name="改名後企業股份有限公司",
            )

    def test_deal_pipeline_progression(self):
        contact = self.crm_service.create_contact(
            company_name="聯發資通企業",
            contact_name="陳處長",
            tax_id="87654321",
        )
        deal = self.crm_service.create_deal(
            contact_id=contact.contact_id,
            title="企業資安系統採購案",
            expected_amount=Decimal("150000.0"),
            stage=DealStage.LEAD,
        )
        self.assertEqual(deal.stage, DealStage.LEAD)

        # Advance: LEAD -> DISCOVERY
        d1 = self.crm_service.advance_deal_stage(deal.deal_id, DealStage.DISCOVERY)
        self.assertEqual(d1.stage, DealStage.DISCOVERY)

        # Advance: DISCOVERY -> PROPOSAL
        d2 = self.crm_service.advance_deal_stage(deal.deal_id, DealStage.PROPOSAL)
        self.assertEqual(d2.stage, DealStage.PROPOSAL)

        # List deals filtered by stage
        proposal_deals = self.crm_service.list_deals(stage=DealStage.PROPOSAL)
        self.assertEqual(len(proposal_deals), 1)
        self.assertEqual(proposal_deals[0].deal_id, deal.deal_id)

    def test_quote_creation_and_auto_stage_update(self):
        contact = self.crm_service.create_contact(
            company_name="台大創新育成科技",
            contact_name="張博士",
            tax_id="11223344",
        )
        deal = self.crm_service.create_deal(
            contact_id=contact.contact_id,
            title="實驗室智慧終端升級專案",
            stage=DealStage.LEAD,
        )

        items = [
            ("SKU-CRM-01", Decimal("2.0"), Decimal("20000.0")),  # 40000
            ("SKU-CRM-02", Decimal("4.0"), Decimal("3500.0")),   # 14000
        ]
        quote = self.crm_service.create_quote(deal.deal_id, items, valid_days=14)
        self.assertEqual(quote.total_amount, Decimal("54000.0"))
        self.assertEqual(len(quote.items), 2)
        self.assertEqual(quote.status, QuoteStatus.DRAFT)

        # Deal should be automatically updated to PROPOSAL and expected_amount updated to 54000
        updated_deal = self.crm_service.get_deal(deal.deal_id)
        self.assertEqual(updated_deal.stage, DealStage.PROPOSAL)
        self.assertEqual(updated_deal.expected_amount, Decimal("54000.0"))

    def test_document_generation_docx_and_pdf(self):
        contact = self.crm_service.create_contact(
            company_name="富邦數位媒體科技",
            contact_name="黃總監",
            tax_id="55667788",
            email="huang@fubon.tw",
            phone="02-87654321",
        )
        deal = self.crm_service.create_deal(
            contact_id=contact.contact_id,
            title="旗艦影音伺服器佈署合約",
            stage=DealStage.PROPOSAL,
        )
        items = [
            ("SKU-CRM-01", Decimal("5.0"), Decimal("19500.0")),
        ]
        quote = self.crm_service.create_quote(deal.deal_id, items)

        # 1. Test DOCX generation
        docx_bytes = self.crm_service.generate_quote_docx(quote.quote_id)
        self.assertGreater(len(docx_bytes), 5000)
        # Verify valid docx by reading with python-docx
        doc = docx.Document(io.BytesIO(docx_bytes))
        doc_text = " ".join(p.text for p in doc.paragraphs)
        self.assertIn("3Dollar ERP", doc_text)
        self.assertIn("富邦數位媒體科技", doc.tables[0].cell(0, 0).text)
        # Verify Table Grid style for clear gridlines
        self.assertEqual(doc.tables[0].style.name, "Table Grid")
        self.assertEqual(doc.tables[1].style.name, "Table Grid")

        # 2. Test PDF generation
        pdf_bytes = self.crm_service.generate_quote_pdf(quote.quote_id)
        self.assertGreater(len(pdf_bytes), 1000)
        # Verify with PyMuPDF
        pdf_doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        self.assertEqual(len(pdf_doc), 1)
        pdf_text = pdf_doc[0].get_text()
        self.assertIn("3Dollar ERP", pdf_text)
        self.assertIn(quote.quote_number, pdf_text)

    def test_convert_deal_to_erp_order(self):
        contact = self.crm_service.create_contact(
            company_name="國泰創新投資",
            contact_name="周協理",
            tax_id="99887766",
        )
        deal = self.crm_service.create_deal(
            contact_id=contact.contact_id,
            title="辦公室自動化核心節點專案",
            stage=DealStage.PROPOSAL,
        )
        items = [
            ("SKU-CRM-01", Decimal("3.0"), Decimal("20000.0")),
        ]
        quote = self.crm_service.create_quote(deal.deal_id, items)

        # Convert to ERP order
        so, won_deal = self.crm_service.convert_deal_to_erp_order(deal.deal_id)

        # Verify SalesOrder created in ERP
        self.assertTrue(so.order_id.startswith("SO-"))
        self.assertEqual(so.customer_id, contact.contact_id)
        self.assertEqual(so.channel, OrderChannel.CRM)
        self.assertEqual(so.status, OrderStatus.DRAFT)
        self.assertEqual(so.total_amount, Decimal("60000.0"))
        self.assertEqual(len(so.items), 1)

        # Verify Deal state updated to WON and linked
        self.assertEqual(won_deal.stage, DealStage.WON)
        self.assertEqual(won_deal.converted_order_id, so.order_id)

        # Verify Quote updated to ACCEPTED
        updated_quote = self.crm_service.get_quote(quote.quote_id)
        self.assertEqual(updated_quote.status, QuoteStatus.ACCEPTED)

        # Attempt to convert again should fail
        with self.assertRaises(ValueError):
            self.crm_service.convert_deal_to_erp_order(deal.deal_id)

        # Attempt to change stage of converted deal should fail
        with self.assertRaises(ValueError):
            self.crm_service.advance_deal_stage(deal.deal_id, DealStage.LOST)


if __name__ == "__main__":
    unittest.main()
