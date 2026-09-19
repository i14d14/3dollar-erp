from __future__ import annotations
import io
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import BinaryIO, Dict, List, Optional, Union
import openpyxl

from src.ai_wrapper.bom.models import BomItem, BomParseResult


class BomExcelParser:
    """Parser for reading Bill of Materials (BOM) Excel files using openpyxl."""

    HEADER_MAPPINGS = {
        "sku": ["料號", "sku", "品號", "材料編號", "零件編號", "part_no", "part number", "item_code"],
        "name": ["品名", "名稱", "物料名稱", "零件名稱", "description", "name", "part name", "item_name"],
        "spec": ["規格", "規格描述", "規格型號", "spec", "specification"],
        "qty": ["數量", "需求數量", "用量", "單位用量", "qty", "quantity", "amount"],
        "unit": ["單位", "unit", "計量單位"],
        "level": ["階層", "level", "階", "層級"],
    }

    def parse(self, file_source: Union[bytes, BinaryIO, str, Path]) -> BomParseResult:
        """Parse BOM Excel file into structured BomParseResult."""
        if isinstance(file_source, bytes):
            wb = openpyxl.load_workbook(io.BytesIO(file_source), data_only=True)
        elif hasattr(file_source, "read"):
            wb = openpyxl.load_workbook(file_source, data_only=True)
        else:
            wb = openpyxl.load_workbook(str(file_source), data_only=True)

        sheet = wb.active
        if not sheet:
            return BomParseResult(errors=["Excel 檔案中找不到有效的工作表 (Worksheet)"])

        # 1. Detect header row (scan first 10 rows)
        header_row_idx = None
        col_map: Dict[str, int] = {}

        for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_idx > 10:
                break
            # Check if this row looks like a header
            temp_map: Dict[str, int] = {}
            for col_idx, cell_value in enumerate(row):
                if cell_value is None:
                    continue
                clean_val = str(cell_value).strip().lower().replace(" ", "")
                for field, aliases in self.HEADER_MAPPINGS.items():
                    if any(clean_val == alias.lower().replace(" ", "") for alias in aliases):
                        temp_map[field] = col_idx
                        break

            # If we found at least sku and qty (or sku and name), consider this the header row
            if "sku" in temp_map and ("qty" in temp_map or "name" in temp_map):
                header_row_idx = row_idx
                col_map = temp_map
                break

        if not header_row_idx:
            return BomParseResult(
                errors=["無法自動辨識 BOM 欄位表頭，請確認試算表中包含「料號/SKU」與「數量/需求量」欄位"]
            )

        items: List[BomItem] = []
        errors: List[str] = []
        total_rows = 0

        # 2. Iterate data rows after header
        for row_idx, row in enumerate(
            sheet.iter_rows(min_row=header_row_idx + 1, values_only=True),
            start=header_row_idx + 1,
        ):
            # Check if row is completely empty
            if not any(row):
                continue

            total_rows += 1
            sku_val = row[col_map["sku"]] if "sku" in col_map and col_map["sku"] < len(row) else None
            if not sku_val or not str(sku_val).strip():
                continue

            sku = str(sku_val).strip()

            # Name
            name_val = row[col_map["name"]] if "name" in col_map and col_map["name"] < len(row) else None
            name = str(name_val).strip() if name_val else sku

            # Spec
            spec_val = row[col_map["spec"]] if "spec" in col_map and col_map["spec"] < len(row) else None
            spec = str(spec_val).strip() if spec_val else None

            # Qty
            qty_val = row[col_map["qty"]] if "qty" in col_map and col_map["qty"] < len(row) else None
            if qty_val is None:
                errors.append(f"第 {row_idx} 列：品號 {sku} 未填寫需求數量")
                continue

            try:
                # Handle possible string formatting or float conversions
                clean_qty_str = str(qty_val).replace(",", "").strip()
                qty = Decimal(clean_qty_str)
                if qty <= 0:
                    errors.append(f"第 {row_idx} 列：品號 {sku} 需求數量必須大於 0（目前為 {qty_val}）")
                    continue
            except (InvalidOperation, ValueError):
                errors.append(f"第 {row_idx} 列：品號 {sku} 需求數量格式無效：'{qty_val}'")
                continue

            # Unit
            unit_val = row[col_map["unit"]] if "unit" in col_map and col_map["unit"] < len(row) else None
            unit = str(unit_val).strip() if unit_val else "PCS"

            # Level
            level_val = row[col_map["level"]] if "level" in col_map and col_map["level"] < len(row) else None
            level = 1
            if level_val is not None:
                try:
                    level = int(str(level_val).strip())
                except ValueError:
                    level = 1

            items.append(
                BomItem(
                    sku=sku,
                    name=name,
                    spec=spec,
                    qty=qty,
                    unit=unit,
                    level=level,
                )
            )

        return BomParseResult(items=items, total_rows=total_rows, errors=errors)
