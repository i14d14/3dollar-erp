from __future__ import annotations
import io
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import BinaryIO, Dict, List, Optional, Union
import openpyxl

from src.core.domain.bom import BomItem, BomParseResult


class BomExcelParser:
    """Parser for reading Bill of Materials (BOM) Excel files using openpyxl."""

    HEADER_MAPPINGS = {
        "sku": ["料號", "sku", "品號", "材料編號", "零件編號", "part_no", "part number", "item_code"],
        "name": ["品名", "名稱", "物料名稱", "零件名稱", "description", "name", "part name", "item_name"],
        "spec": ["規格", "規格描述", "規格型號", "spec", "specification"],
        "qty": ["數量", "需求數量", "用量", "單位用量", "qty", "quantity", "amount"],
        "unit": ["單位", "unit", "計量單位"],
        "level": ["階層", "level", "階", "層級"],
        "supplier": ["供應商", "廠商", "供應商名稱", "廠商名稱", "supplier", "vendor"],
        "unit_cost": ["單價", "採購單價", "成本", "預估單價", "進價", "price", "unit_price", "cost", "unit_cost"],
    }

    def parse(self, file_source: Union[bytes, BinaryIO, str, Path]) -> BomParseResult:
        """Parse BOM Excel file into structured BomParseResult."""
        if isinstance(file_source, bytes):
            wb = openpyxl.load_workbook(io.BytesIO(file_source), data_only=True)
        elif hasattr(file_source, "read"):
            wb = openpyxl.load_workbook(file_source, data_only=True)
        else:
            wb = openpyxl.load_workbook(str(file_source), data_only=True)

        sheet = wb.active
        if not sheet:
            return BomParseResult(errors=["Excel 檔案中找不到有效的工作表 (Worksheet)"])

        # 1. Detect header row (scan first 10 rows)
        header_row_idx = None
        col_map: Dict[str, int] = {}

        for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_idx > 10:
                break
            # Check if this row looks like a header
            temp_map: Dict[str, int] = {}
            for col_idx, cell_value in enumerate(row):
                if cell_value is None:
                    continue
                clean_val = str(cell_value).strip().lower().replace(" ", "")
                for field, aliases in self.HEADER_MAPPINGS.items():
                    if any(clean_val == alias.lower().replace(" ", "") for alias in aliases):
                        temp_map[field] = col_idx
                        break

            # If we found at least sku and qty (or sku and name), consider this the header row
            if "sku" in temp_map and ("qty" in temp_map or "name" in temp_map):
                header_row_idx = row_idx
                col_map = temp_map
                break

        if not header_row_idx:
            return BomParseResult(
                errors=["無法自動辨識 BOM 欄位表頭。請確保檔案包含「料號/SKU」與「需求數量/用量」欄位。"]
            )

        items: List[BomItem] = []
        errors: List[str] = []
        total_rows = 0

        # 2. Iterate data rows after header
        for row_idx, row in enumerate(
            sheet.iter_rows(min_row=header_row_idx + 1, values_only=True),
            start=header_row_idx + 1,
        ):
            # Check if entire row is empty
            if not any(row):
                continue

            total_rows += 1

            # Extract fields using col_map
            raw_sku = row[col_map["sku"]] if "sku" in col_map and col_map["sku"] < len(row) else None
            if not raw_sku:
                continue

            sku = str(raw_sku).strip().upper()

            raw_name = row[col_map["name"]] if "name" in col_map and col_map["name"] < len(row) else sku
            name = str(raw_name).strip() if raw_name is not None else sku

            raw_spec = row[col_map["spec"]] if "spec" in col_map and col_map["spec"] < len(row) else None
            spec = str(raw_spec).strip() if raw_spec is not None else None

            raw_unit = row[col_map["unit"]] if "unit" in col_map and col_map["unit"] < len(row) else "PCS"
            unit = str(raw_unit).strip().upper() if raw_unit is not None else "PCS"

            raw_level = row[col_map["level"]] if "level" in col_map and col_map["level"] < len(row) else 1
            try:
                level = int(raw_level) if raw_level is not None else 1
            except (ValueError, TypeError):
                level = 1

            raw_qty = row[col_map["qty"]] if "qty" in col_map and col_map["qty"] < len(row) else None
            if raw_qty is None:
                errors.append(f"第 {row_idx} 行 (料號: {sku}) 缺少數量，略過該列。")
                continue

            try:
                qty_clean = str(raw_qty).strip().replace(",", "")
                qty = Decimal(qty_clean)
                if qty <= Decimal("0"):
                    errors.append(f"第 {row_idx} 行 (料號: {sku}) 數量必須大於 0 (填入: {raw_qty})。")
                    continue
            except (InvalidOperation, ValueError):
                errors.append(f"第 {row_idx} 行 (料號: {sku}) 數量數值格式錯誤 (填入: {raw_qty})。")
                continue

            raw_supplier = row[col_map["supplier"]] if "supplier" in col_map and col_map["supplier"] < len(row) else None
            supplier_name = str(raw_supplier).strip() if raw_supplier is not None and str(raw_supplier).strip() else None

            raw_cost = row[col_map["unit_cost"]] if "unit_cost" in col_map and col_map["unit_cost"] < len(row) else None
            unit_cost = None
            if raw_cost is not None and str(raw_cost).strip():
                try:
                    c_clean = str(raw_cost).strip().replace(",", "").replace("$", "")
                    c_dec = Decimal(c_clean)
                    if c_dec > Decimal("0"):
                        unit_cost = c_dec
                except (InvalidOperation, ValueError):
                    pass

            items.append(
                BomItem(
                    sku=sku,
                    name=name,
                    spec=spec,
                    qty=qty,
                    unit=unit,
                    level=level,
                    supplier_name=supplier_name,
                    unit_cost=unit_cost,
                )
            )

        return BomParseResult(
            items=items,
            total_rows=total_rows,
            errors=errors,
        )
